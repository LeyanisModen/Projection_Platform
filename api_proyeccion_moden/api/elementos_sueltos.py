"""Excel de elementos sueltos de un proyecto (barras, vigas, zunchos...).

Son piezas que van a obra fuera de los modulos, antes o despues que ellos.
La oficina sube el Excel tal cual; el dashboard de la ferralla lo ve como una
tabla por hoja y puede descargar el original. Aqui solo se lee: no se asume
ningun formato de columnas. La cabecera es la primera fila con al menos dos
celdas; lo que haya encima (titulos de una sola celda) se devuelve como notas.
"""
import os
from datetime import date, datetime, time

# Extensiones que se aceptan al subir. Los .xls antiguos se guardan y se
# pueden descargar, pero openpyxl no los lee, asi que no se previsualizan.
EXTENSIONES = ('.xlsx', '.xlsm', '.xls')
PREVISUALIZABLES = ('.xlsx', '.xlsm')
# Tope de filas por hoja que viajan al navegador; el resto, en el Excel.
MAX_FILAS = 2000


def _valor(celda):
    """Valor de celda listo para JSON: numeros como numeros, fechas en texto."""
    if celda is None:
        return None
    if isinstance(celda, bool):
        return 'Sí' if celda else 'No'
    if isinstance(celda, datetime):
        if celda.time() == time(0, 0):
            return celda.strftime('%d/%m/%Y')
        return celda.strftime('%d/%m/%Y %H:%M')
    if isinstance(celda, date):
        return celda.strftime('%d/%m/%Y')
    if isinstance(celda, time):
        return celda.strftime('%H:%M')
    if isinstance(celda, float):
        if celda.is_integer():
            return int(celda)
        return round(celda, 4)
    if isinstance(celda, int):
        return celda
    texto = str(celda).strip()
    return texto or None


# Filas de titulo que se admiten encima de la cabecera.
MAX_NOTAS = 10


def _rellenas(fila):
    return sum(1 for valor in fila if valor is not None)


def _leer_hoja(hoja):
    filas = []
    total = 0
    for fila in hoja.iter_rows(values_only=True):
        valores = [_valor(celda) for celda in fila]
        if not _rellenas(valores):
            continue
        total += 1
        # Algo mas que el tope: notas de cabecera y una fila para saber que se recorta.
        if len(filas) <= MAX_FILAS + MAX_NOTAS:
            filas.append(valores)
    if not filas:
        return None

    # La cabecera es la primera fila con dos o mas celdas; si ninguna tiene
    # dos (hoja de una sola columna), la primera.
    inicio = next(
        (indice for indice, fila in enumerate(filas[:MAX_NOTAS + 1]) if _rellenas(fila) >= 2),
        0,
    )
    notas = [
        ' '.join(str(valor) for valor in fila if valor is not None)
        for fila in filas[:inicio]
    ]
    cabecera, datos = filas[inicio], filas[inicio + 1:]
    total -= inicio + 1
    tabla = [cabecera, *datos]
    ancho = max(len(fila) for fila in tabla)
    # Fuera las columnas sin nada, ni cabecera ni datos.
    usadas = [
        indice for indice in range(ancho)
        if any(indice < len(fila) and fila[indice] is not None for fila in tabla)
    ]
    columnas = []
    for posicion, indice in enumerate(usadas, start=1):
        titulo = cabecera[indice] if indice < len(cabecera) else None
        columnas.append(str(titulo) if titulo is not None else f'Columna {posicion}')
    datos = [
        [fila[indice] if indice < len(fila) else None for indice in usadas]
        for fila in datos[:MAX_FILAS]
    ]
    return {
        'nombre': hoja.title,
        'notas': notas,
        'columnas': columnas,
        'filas': datos,
        'total_filas': total,
        'recortado': total > len(datos),
    }


def leer_elementos_sueltos(field_file):
    """Hojas visibles del Excel como ``{nombre, columnas, filas, ...}``.

    Nunca lanza: si el fichero no se puede leer, ``previsualizable`` es False
    y ``motivo`` lo explica, para que el dashboard ofrezca la descarga.
    """
    nombre = os.path.basename(field_file.name)
    extension = os.path.splitext(nombre)[1].lower()
    resultado = {
        'nombre_archivo': nombre,
        'previsualizable': False,
        'motivo': None,
        'hojas': [],
    }
    if extension not in PREVISUALIZABLES:
        resultado['motivo'] = (
            'Este Excel está en formato antiguo (.xls) y no se puede ver aquí. '
            'Descárgalo para abrirlo.'
        )
        return resultado

    from openpyxl import load_workbook

    libro = None
    try:
        field_file.open('rb')
        libro = load_workbook(field_file, read_only=True, data_only=True)
        for hoja in libro.worksheets:
            if getattr(hoja, 'sheet_state', 'visible') != 'visible':
                continue
            contenido = _leer_hoja(hoja)
            if contenido is not None:
                resultado['hojas'].append(contenido)
        resultado['previsualizable'] = True
    except FileNotFoundError:
        resultado['motivo'] = 'El archivo no está disponible en el servidor.'
    except Exception:
        resultado['motivo'] = 'No se ha podido leer el Excel. Descárgalo para abrirlo.'
    finally:
        if libro is not None:
            libro.close()
        try:
            field_file.close()
        except Exception:
            pass
    return resultado
