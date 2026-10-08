"""Archivar un proyecto: todo lo suyo en un .zip y, despues, liberar espacio.

Archivar no borra el proyecto. Borra lo que ocupa (presentaciones, fotos,
documentos y adjuntos de la lista de control) y deja los datos ligeros
(modulos, bastidores, pesos, fechas, tiempos, colas hechas, lista de control)
para que las estadisticas de la ferralla sigan completas.

El .zip se genera mientras se descarga (sin fichero temporal ni todo en
memoria): las imagenes van sin comprimir, que ya lo estan, y el Excel, el
JSON y el LEEME con deflate.
"""
import io
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path

from django.conf import settings
from django.core import signing
from django.core.serializers.json import DjangoJSONEncoder
from django.db import transaction
from django.utils import timezone

from .models import (
    DetalleModuloFase, EventoFabricacion, FotoFabricacion, GrupoBastidor, GrupoMesasProyecto,
    Imagen, MaterialInformado, MaterialPieza, MesaQueueItem, MesaQueueStatus, Modulo,
    ProyectoCheck, ProyectoCheckAdjunto,
)

SALT = 'archivo-proyecto'
# El enlace de descarga caduca: se pide justo antes de descargar.
TOKEN_MAX_AGE = 30 * 60
CHUNK = 1024 * 1024
FASE_CARPETA = {'INFERIOR': 'INF', 'SUPERIOR': 'SUP'}
CAMPOS_FICHERO = ('plano_archivo', 'documentos_archivo', 'elementos_sueltos_archivo', 'fichero_datos_tecnicos')


# ---------------------------------------------------------------------------
# Antes de archivar
# ---------------------------------------------------------------------------
def motivos_para_no_archivar(proyecto):
    """Por que no se puede archivar todavia (lista vacia: se puede)."""
    motivos = []
    if proyecto.archivado_at:
        motivos.append('El proyecto ya está archivado.')
    if MesaQueueItem.objects.filter(
        modulo__proyecto=proyecto,
        status__in=[MesaQueueStatus.EN_COLA, MesaQueueStatus.MOSTRANDO],
    ).exists():
        motivos.append('Tiene módulos en las mesas: quítalo de producción antes de archivarlo.')
    if GrupoMesasProyecto.objects.filter(proyecto=proyecto).exists():
        motivos.append('Está en la cola de una línea: quítalo desde Gestionar antes de archivarlo.')
    return motivos


def token_descarga(proyecto, usuario):
    return signing.dumps({'p': proyecto.pk, 'u': usuario.pk}, salt=SALT)


def leer_token(token):
    """(proyecto_id, usuario_id) o BadSignature/SignatureExpired."""
    datos = signing.loads(token, salt=SALT, max_age=TOKEN_MAX_AGE)
    return datos['p'], datos['u']


def nombre_zip(proyecto):
    base = re.sub(r'[^A-Za-z0-9._-]+', '_', proyecto.nombre).strip('_') or f'proyecto_{proyecto.pk}'
    return f'{base}_archivo_{timezone.localdate():%Y%m%d}.zip'


# ---------------------------------------------------------------------------
# Contenido
# ---------------------------------------------------------------------------
def _ruta_media(valor):
    """Fichero del volumen a partir de una URL /media/... o de un nombre de FileField."""
    if not valor:
        return None
    relativo = str(valor)
    prefijo = settings.MEDIA_URL or '/media/'
    if relativo.startswith(prefijo):
        relativo = relativo[len(prefijo):]
    raiz = Path(settings.MEDIA_ROOT).resolve()
    ruta = (raiz / relativo.lstrip('/')).resolve()
    return ruta if raiz in ruta.parents else None


def _parte(texto):
    """Trozo de ruta dentro del zip, sin barras ni caracteres raros."""
    return re.sub(r'[\\/:*?"<>|]+', '_', str(texto or '')).strip() or 'sin_nombre'


def _local(valor):
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return timezone.localtime(valor).replace(tzinfo=None) if timezone.is_aware(valor) else valor
    return valor


def _nombre_usuario(usuario):
    return (usuario.get_full_name() or usuario.username) if usuario else None


def _datos(proyecto):
    """Filas de todo lo que tiene el proyecto, para el Excel y el JSON."""
    modulos = list(
        Modulo.objects.filter(proyecto=proyecto)
        .select_related('grupo_bastidor')
        .order_by('grupo_bastidor__indice', 'orden_intra', 'nombre')
    )
    nombres = {m.pk: m.nombre for m in modulos}
    campos_detalle = [
        f.name for f in DetalleModuloFase._meta.concrete_fields
        if f.name not in ('id', 'modulo', 'fase', 'created_at', 'updated_at')
    ]
    detalles = {}
    for d in DetalleModuloFase.objects.filter(modulo__proyecto=proyecto):
        detalles[(d.modulo_id, d.fase)] = d

    filas_modulos = []
    for m in modulos:
        fila = {
            'Módulo': m.nombre,
            'Bastidor': m.grupo_bastidor.nombre if m.grupo_bastidor else None,
            'Orden': m.orden_intra,
            'Tipo': m.tipo_modulo,
            'Ancho (cm)': m.ancho_cm,
            'Estado': m.estado,
            'Inferior hecho': _local(m.inferior_completado_at),
            'Superior hecho': _local(m.superior_completado_at),
            'Completado': _local(m.completado_at),
            'Cerrado': _local(m.cerrado_at) if m.cerrado else None,
            'Códigos de color': m.codigos_color,
        }
        for fase, corto in FASE_CARPETA.items():
            detalle = detalles.get((m.pk, fase))
            for campo in campos_detalle:
                fila[f'{corto} {campo}'] = getattr(detalle, campo) if detalle else None
        filas_modulos.append(fila)

    bastidores = [
        {'Bastidor': g.nombre, 'Índice': g.indice, 'Sufijo': g.sufijo, 'Módulos': g.modulos.count()}
        for g in GrupoBastidor.objects.filter(proyecto=proyecto).order_by('indice', 'sufijo')
    ]

    fotos = []
    for foto in FotoFabricacion.objects.filter(modulo__proyecto=proyecto).select_related('modulo', 'mesa').order_by('modulo_id', 'fase', 'paso', 'capturada_at'):
        fotos.append({
            'Módulo': foto.modulo.nombre,
            'Fase': foto.fase,
            'Paso': foto.paso + 1,
            'Mesa': foto.mesa.nombre if foto.mesa else None,
            'Capturada': _local(foto.capturada_at),
            'Comprobación': {True: 'Correcta', False: 'Fallida'}.get(foto.check_result, 'Sin comprobar'),
            'Detalle comprobación': json.dumps(foto.check_detail, ensure_ascii=False) if foto.check_detail else None,
            'Fichero': None,  # se rellena al meterla en el zip
            '_foto': foto,
        })

    checks = []
    for paso in ProyectoCheck.objects.filter(proyecto=proyecto).select_related('completado_por').prefetch_related('adjuntos').order_by('orden', 'id'):
        checks.append({
            'Orden': paso.orden,
            'Paso': paso.titulo,
            'Fecha límite': paso.fecha_limite,
            'Días antes del montaje': paso.dias_antes_montaje,
            'Completado': 'Sí' if paso.completado else 'No',
            'Completado el': _local(paso.completado_at),
            'Completado por': _nombre_usuario(paso.completado_por),
            'Requiere documento': 'Sí' if paso.requiere_documento else 'No',
            'Bloquea producción': 'Sí' if paso.bloquea_produccion else 'No',
            'Documentos': ', '.join(a.nombre_original for a in paso.adjuntos.all()) or None,
        })

    eventos = [
        {
            'Módulo': nombres.get(e.modulo_id), 'Fase': e.fase, 'Mesa': e.mesa.nombre if e.mesa else None,
            'Tipo': e.tipo, 'Paso': (e.paso + 1) if e.paso is not None else None, 'Fecha': _local(e.at),
        }
        for e in EventoFabricacion.objects.filter(modulo__proyecto=proyecto).select_related('mesa').order_by('at', 'id')
    ]
    colas = [
        {
            'Módulo': nombres.get(i.modulo_id), 'Fase': i.fase, 'Mesa': i.mesa.nombre if i.mesa else None,
            'Estado': i.status, 'Asignado el': _local(i.assigned_at), 'Hecho el': _local(i.done_at),
            'Hecho por': _nombre_usuario(i.done_by),
        }
        for i in MesaQueueItem.objects.filter(modulo__proyecto=proyecto).select_related('mesa', 'done_by').order_by('mesa_id', 'position', 'id')
    ]
    materiales = [
        {'Módulo': nombres.get(p.modulo_id), 'Tipo': p.tipo, 'Capa': p.capa, 'Subtipo': p.subtipo, 'Longitud': p.longitud}
        for p in MaterialPieza.objects.filter(proyecto=proyecto).order_by('modulo_id', 'id')
    ]
    informados = [
        {'Material': i.clave_material, 'Informado': 'Sí' if i.informado else 'No', 'Origen': i.origen, 'Marcado el': _local(i.fecha_marcado)}
        for i in MaterialInformado.objects.filter(proyecto=proyecto).order_by('clave_material')
    ]
    return {
        'modulos': filas_modulos, 'bastidores': bastidores, 'fotos': fotos, 'lista_control': checks,
        'eventos': eventos, 'colas': colas, 'materiales': materiales, 'materiales_informados': informados,
    }


def _estadisticas(proyecto, usuario):
    """Las mismas estadisticas que la pagina, de todo el proyecto."""
    from rest_framework.test import APIRequestFactory, force_authenticate

    from .views import ProductionStatsView

    peticion = APIRequestFactory().get('/api/stats/production/', {'proyecto': proyecto.pk, 'rango': 'proyecto'})
    force_authenticate(peticion, user=usuario)
    respuesta = ProductionStatsView.as_view()(peticion)
    return respuesta.data if respuesta.status_code == 200 else None


def _plano(valor, prefijo=''):
    """Dict anidado -> filas clave/valor para una hoja de resumen."""
    filas = []
    for clave, v in (valor or {}).items():
        nombre = f'{prefijo}{clave}'
        if isinstance(v, dict):
            filas.extend(_plano(v, f'{nombre}.'))
        elif isinstance(v, list):
            filas.append((nombre, json.dumps(v, ensure_ascii=False, cls=DjangoJSONEncoder)))
        else:
            filas.append((nombre, v))
    return filas


def _celda(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False, cls=DjangoJSONEncoder)
    if isinstance(v, datetime) and timezone.is_aware(v):
        return _local(v)
    return v


def _hoja(libro, titulo, filas):
    hoja = libro.create_sheet(titulo[:31])
    filas = [{k: v for k, v in fila.items() if not k.startswith('_')} for fila in filas]
    columnas = []
    for fila in filas:
        for clave in fila:
            if clave not in columnas:
                columnas.append(clave)
    if not columnas:
        hoja.append(['Sin datos'])
        return
    hoja.append(columnas)
    for fila in filas:
        hoja.append([_celda(fila.get(c)) for c in columnas])


def _excel(proyecto, datos, stats):
    from openpyxl import Workbook

    libro = Workbook()
    portada = libro.active
    portada.title = 'Proyecto'
    totales = (stats or {}).get('totals') or {}
    resumen = [
        ('Proyecto', proyecto.nombre),
        ('Ferralla', _nombre_usuario(proyecto.usuario)),
        ('Fecha de montaje', proyecto.fecha_montaje),
        ('Módulos', len(datos['modulos'])),
        ('Módulos completados', sum(1 for m in datos['modulos'] if m['Completado'])),
        ('Bastidores', len(datos['bastidores'])),
        ('Fotos de fabricación', len(datos['fotos'])),
        ('Peso total fabricado (kg)', totals_peso(totales)),
        ('Desperdicio (kg)', totales.get('desperdicio_kg')),
        ('Estadísticas desde', (stats or {}).get('range', {}).get('from')),
        ('Estadísticas hasta', (stats or {}).get('range', {}).get('to')),
        ('Archivo generado', _local(timezone.now())),
    ]
    portada.append(['Campo', 'Valor'])
    for fila in resumen:
        portada.append([fila[0], _celda(fila[1])])

    _hoja(libro, 'Módulos', datos['modulos'])
    _hoja(libro, 'Bastidores', datos['bastidores'])
    if stats:
        _hoja(libro, 'Estadísticas por módulo', stats.get('modulos') or [])
        _hoja(libro, 'Estadísticas por día', stats.get('por_dia') or [])
        _hoja(libro, 'Estadísticas por mesa', stats.get('por_mesa') or [])
        resumen_stats = libro.create_sheet('Resumen estadísticas')
        resumen_stats.append(['Dato', 'Valor'])
        for clave in ('totals', 'tiempos', 'esperado', 'proyecto', 'range'):
            for nombre, valor in _plano(stats.get(clave), f'{clave}.'):
                resumen_stats.append([nombre, _celda(valor)])
    _hoja(libro, 'Fotos', datos['fotos'])
    _hoja(libro, 'Lista de control', datos['lista_control'])
    _hoja(libro, 'Fabricación (eventos)', datos['eventos'])
    _hoja(libro, 'Fabricación (mesas)', datos['colas'])
    _hoja(libro, 'Materiales', datos['materiales'])
    _hoja(libro, 'Materiales informados', datos['materiales_informados'])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def totals_peso(totales):
    return totales.get('peso_total_kg')


LEEME = """ARCHIVO DEL PROYECTO {nombre}
Generado el {fecha} desde la plataforma de proyección de Moden.

datos.xlsx
    Todo lo del proyecto en hojas: resumen, módulos (bastidor, estado, fechas
    de cada fase, pesos, malla, desperdicio, dificultad), bastidores,
    estadísticas de todo el proyecto (por módulo, por día y por mesa, las
    mismas de la página de Estadísticas), fotos, lista de control, eventos de
    fabricación (inicio, pasos, pausas, fin), mesas y colas, y materiales.

proyecto.json
    Lo mismo en formato JSON, para leerlo con otros programas.

presentaciones/MÓDULO/INF|SUP/PLAYER y MONITOR
    Las imágenes que se proyectaron (PLAYER) y las de la pantalla del
    operario (MONITOR), con sus nombres originales y la misma estructura
    de carpetas que se usa para importar.

fotos/MÓDULO/FASE
    Fotos de fabricación. El nombre empieza por el paso y la fecha; en la
    hoja Fotos de datos.xlsx está cada una con su mesa y el resultado de
    la comprobación de colores.

documentos/
    Plano, documentación, elementos sueltos y base técnica (.db) originales.

lista_control/
    Documentos adjuntos a cada paso de la lista de control.

Al archivar, la plataforma borra estos ficheros pero conserva los datos de
fabricación (módulos, pesos, fechas y tiempos): las estadísticas de la
ferralla siguen contando este proyecto.
"""


def entradas(proyecto, usuario):
    """(ruta en el zip, Path del volumen | bytes | str), en el orden en que se escriben."""
    yield 'LEEME.txt', LEEME.format(nombre=proyecto.nombre, fecha=timezone.localtime(timezone.now()).strftime('%d/%m/%Y %H:%M'))
    datos = _datos(proyecto)
    try:
        stats = _estadisticas(proyecto, usuario)
    except Exception:
        stats = None

    # Ficheros: presentaciones, fotos, documentos y adjuntos.
    ficheros = []
    for img in Imagen.objects.filter(modulo__proyecto=proyecto, activo=True).select_related('modulo').order_by('modulo_id', 'fase', 'orden', 'id'):
        fase = FASE_CARPETA.get(img.fase, _parte(img.fase))
        for pantalla, url in (('PLAYER', img.url or (img.archivo.name if img.archivo else None)), ('MONITOR', img.url_monitor)):
            ruta = _ruta_media(url)
            if ruta:
                ficheros.append((f'presentaciones/{_parte(img.modulo.nombre)}/{fase}/{pantalla}/{_parte(ruta.name)}', ruta))
    for fila in datos['fotos']:
        foto = fila['_foto']
        ruta = _ruta_media(foto.url)
        if ruta:
            momento = _local(foto.capturada_at)
            nombre = f'fotos/{_parte(foto.modulo.nombre)}/{FASE_CARPETA.get(foto.fase, _parte(foto.fase))}/' \
                     f'{foto.paso + 1:02d}_{momento:%Y%m%d-%H%M%S}_{_parte(ruta.name)}'
            fila['Fichero'] = nombre
            ficheros.append((nombre, ruta))
    for campo in CAMPOS_FICHERO:
        valor = getattr(proyecto, campo)
        ruta = _ruta_media(valor.name) if valor else None
        if ruta:
            ficheros.append((f'documentos/{_parte(ruta.name)}', ruta))
    for adjunto in ProyectoCheckAdjunto.objects.filter(paso__proyecto=proyecto).select_related('paso').order_by('paso__orden', 'id'):
        ruta = _ruta_media(adjunto.archivo.name) if adjunto.archivo else None
        if ruta:
            ficheros.append((f'lista_control/{adjunto.paso.orden:02d}_{_parte(adjunto.paso.titulo)}/{_parte(adjunto.nombre_original or ruta.name)}', ruta))

    yield 'datos.xlsx', _excel(proyecto, datos, stats)
    yield 'proyecto.json', json.dumps({
        'proyecto': {
            'id': proyecto.pk, 'nombre': proyecto.nombre, 'ferralla': _nombre_usuario(proyecto.usuario),
            'fecha_montaje': proyecto.fecha_montaje, 'modulos_previstos': proyecto.modulos_previstos,
            'estrategia_bastidor': proyecto.estrategia_bastidor,
        },
        **{clave: [{k: v for k, v in fila.items() if not k.startswith('_')} for fila in filas] for clave, filas in datos.items()},
        'estadisticas': stats,
    }, ensure_ascii=False, indent=2, cls=DjangoJSONEncoder)
    yield from ficheros


class _Salida:
    """Destino del ZipFile que guarda lo escrito hasta que se envia."""

    def __init__(self):
        self._trozos = []
        self._posicion = 0

    def write(self, datos):
        self._trozos.append(bytes(datos))
        self._posicion += len(datos)
        return len(datos)

    def tell(self):
        return self._posicion

    def flush(self):
        pass

    def vaciar(self):
        datos = b''.join(self._trozos)
        self._trozos = []
        return datos


def zip_en_streaming(proyecto, usuario):
    """Genera el .zip a trozos; si falta algun fichero lo apunta en FALTAN.txt."""
    salida = _Salida()
    usados = set()
    faltan = []
    ahora = timezone.localtime(timezone.now()).timetuple()[:6]
    with zipfile.ZipFile(salida, 'w', allowZip64=True) as archivo:
        for nombre, contenido in entradas(proyecto, usuario):
            base, punto, extension = nombre.rpartition('.')
            unico, n = nombre, 1
            while unico in usados:
                n += 1
                unico = f'{base}_{n}.{extension}' if punto else f'{nombre}_{n}'
            usados.add(unico)
            if isinstance(contenido, Path):
                if not contenido.is_file():
                    faltan.append(unico)
                    continue
                info = zipfile.ZipInfo(unico, date_time=datetime.fromtimestamp(contenido.stat().st_mtime).timetuple()[:6])
                info.compress_type = zipfile.ZIP_STORED
                with contenido.open('rb') as origen, archivo.open(info, 'w') as destino:
                    while trozo := origen.read(CHUNK):
                        destino.write(trozo)
                        yield salida.vaciar()
            else:
                info = zipfile.ZipInfo(unico, date_time=ahora)
                info.compress_type = zipfile.ZIP_DEFLATED
                archivo.writestr(info, contenido.encode('utf-8') if isinstance(contenido, str) else contenido)
            yield salida.vaciar()
        if faltan:
            info = zipfile.ZipInfo('FALTAN.txt', date_time=ahora)
            info.compress_type = zipfile.ZIP_DEFLATED
            archivo.writestr(info, ('Ficheros que no estaban en el servidor:\n' + '\n'.join(faltan)).encode('utf-8'))
    yield salida.vaciar()


# ---------------------------------------------------------------------------
# Liberar espacio
# ---------------------------------------------------------------------------
def archivar(proyecto, usuario):
    """Borra los ficheros del proyecto y lo marca como archivado.

    Se quedan modulos, detalles tecnicos, bastidores, colas hechas, eventos,
    lista de control (sin adjuntos) y materiales: las estadisticas no cambian.
    """
    from .project_media import collect_project_media, delete_project_media

    snapshot = collect_project_media(proyecto)
    with transaction.atomic():
        Imagen.objects.filter(modulo__proyecto=proyecto).delete()
        FotoFabricacion.objects.filter(modulo__proyecto=proyecto).delete()
        ProyectoCheckAdjunto.objects.filter(paso__proyecto=proyecto).delete()
        for campo in CAMPOS_FICHERO:
            setattr(proyecto, campo, None)
        proyecto.archivado_at = timezone.now()
        proyecto.archivado_por = usuario
        proyecto.save(update_fields=[*CAMPOS_FICHERO, 'archivado_at', 'archivado_por'])
        transaction.on_commit(lambda: delete_project_media(snapshot), robust=True)
    return proyecto
