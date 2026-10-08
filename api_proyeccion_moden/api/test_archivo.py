import io
import json
import shutil
import tempfile
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone
from rest_framework.test import APITestCase

from .models import (
    DetalleModuloFase, FotoFabricacion, GrupoBastidor, GrupoMesas, GrupoMesasProyecto, Imagen,
    Mesa, MesaQueueItem, Modulo, Proyecto, ProyectoCheck, ProyectoCheckAdjunto,
)
from .planning import plan_ferralla


class ArchivoProyectoTests(APITestCase):
    def setUp(self):
        self.media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)
        override = override_settings(MEDIA_ROOT=self.media)
        override.enable()
        self.addCleanup(override.disable)

        self.admin = User.objects.create_user('oficina', is_staff=True)
        self.factory = User.objects.create_user('ferralla')
        self.project = Proyecto.objects.create(nombre='Obra Prueba', usuario=self.factory, fecha_montaje=date(2026, 11, 20))
        bastidor = GrupoBastidor.objects.create(proyecto=self.project, indice=1, nombre='Grupo 1')
        self.modulo = Modulo.objects.create(nombre='A01', proyecto=self.project, grupo_bastidor=bastidor, orden_intra=1)
        DetalleModuloFase.objects.create(modulo=self.modulo, fase='INFERIOR', peso_malla_final_kg=Decimal('10.50'))

        pid, mid = self.project.pk, self.modulo.pk
        self.player = self._fichero(f'imagenes/{pid}/{mid}/MOD_A01_INF_00.png', b'player')
        self.monitor = self._fichero(f'imagenes/{pid}/{mid}/MOD_A01_INF_MONITOR_00.png', b'monitor')
        Imagen.objects.create(
            modulo=self.modulo, fase='INFERIOR', orden=1,
            url=f'/media/imagenes/{pid}/{mid}/MOD_A01_INF_00.png',
            url_monitor=f'/media/imagenes/{pid}/{mid}/MOD_A01_INF_MONITOR_00.png',
        )
        self.foto = self._fichero(f'fotos/{pid}/{mid}/captura.jpg', b'foto')
        FotoFabricacion.objects.create(
            modulo=self.modulo, fase='INFERIOR', paso=2, url=f'/media/fotos/{pid}/{mid}/captura.jpg', check_result=False,
        )
        self.project.plano_archivo = SimpleUploadedFile('plano.pdf', b'%PDF plano')
        self.project.save()
        paso = ProyectoCheck.objects.create(proyecto=self.project, titulo='Planos', orden=1)
        ProyectoCheckAdjunto.objects.create(
            paso=paso, archivo=SimpleUploadedFile('correo.eml', b'correo'), nombre_original='correo.eml', tamano=6,
        )

        # Fabricado en una mesa: es lo que cuentan las estadisticas.
        hecho = timezone.make_aware(datetime(2026, 10, 6, 10, 0))
        mesa = Mesa.objects.create(nombre='Mesa 1', usuario=self.factory, tipo='INFERIOR', indice=1)
        MesaQueueItem.objects.create(mesa=mesa, modulo=self.modulo, fase='INFERIOR', status='HECHO', done_at=hecho, position=0)
        Modulo.objects.filter(pk=mid).update(inferior_hecho=True, superior_hecho=True, estado='COMPLETADO', completado_at=hecho)
        self.client.force_authenticate(self.admin)

    def _fichero(self, relativo, contenido):
        ruta = Path(self.media) / relativo
        ruta.parent.mkdir(parents=True, exist_ok=True)
        ruta.write_bytes(contenido)
        return ruta

    def _descargar(self):
        preparado = self.client.post(f'/api/proyectos/{self.project.pk}/archivo/')
        self.assertEqual(preparado.status_code, 200, preparado.data)
        self.assertTrue(preparado.data['nombre'].startswith('Obra_Prueba_archivo_'))
        # El enlace firmado vale sin sesion: el navegador lo abre como descarga.
        self.client.force_authenticate(None)
        respuesta = self.client.get(preparado.data['url'])
        self.client.force_authenticate(self.admin)
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], 'application/zip')
        self.assertEqual(respuesta['X-Accel-Buffering'], 'no')
        return zipfile.ZipFile(io.BytesIO(b''.join(respuesta.streaming_content)))

    def _modulos_completados(self):
        stats = self.client.get(f'/api/stats/production/?proyecto={self.project.pk}&rango=proyecto').data
        return stats['totals']['modulos_completados']

    def test_el_zip_lleva_presentaciones_fotos_documentos_y_datos(self):
        archivo = self._descargar()
        nombres = set(archivo.namelist())
        for nombre in (
            'LEEME.txt', 'datos.xlsx', 'proyecto.json',
            'presentaciones/A01/INF/PLAYER/MOD_A01_INF_00.png',
            'presentaciones/A01/INF/MONITOR/MOD_A01_INF_MONITOR_00.png',
            'documentos/plano.pdf',
            'lista_control/01_Planos/correo.eml',
        ):
            self.assertIn(nombre, nombres)
        self.assertEqual(archivo.read('presentaciones/A01/INF/MONITOR/MOD_A01_INF_MONITOR_00.png'), b'monitor')
        fotos = [n for n in nombres if n.startswith('fotos/A01/INF/03_')]
        self.assertEqual(len(fotos), 1, nombres)
        self.assertEqual(archivo.read(fotos[0]), b'foto')

        datos = json.loads(archivo.read('proyecto.json'))
        self.assertEqual(datos['proyecto']['nombre'], 'Obra Prueba')
        self.assertEqual(datos['modulos'][0]['Módulo'], 'A01')
        self.assertEqual(datos['modulos'][0]['INF peso_malla_final_kg'], '10.50')
        self.assertEqual((datos['fotos'][0]['Comprobación'], datos['fotos'][0]['Fichero']), ('Fallida', fotos[0]))
        self.assertEqual(datos['estadisticas']['totals']['modulos_completados'], 1)

        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(archivo.read('datos.xlsx')))
        for hoja in ('Proyecto', 'Módulos', 'Estadísticas por módulo', 'Estadísticas por día', 'Fotos', 'Lista de control'):
            self.assertIn(hoja, libro.sheetnames)
        self.assertEqual(libro['Módulos']['A2'].value, 'A01')

    def test_archivar_borra_los_ficheros_y_las_estadisticas_siguen(self):
        self.assertEqual(self._modulos_completados(), 1)
        sin_confirmar = self.client.post(f'/api/proyectos/{self.project.pk}/archivar/', {}, format='json')
        self.assertEqual(sin_confirmar.status_code, 400)

        with self.captureOnCommitCallbacks(execute=True):
            respuesta = self.client.post(f'/api/proyectos/{self.project.pk}/archivar/', {'confirmado': True}, format='json')
        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        self.assertIsNotNone(respuesta.data['archivado_at'])
        self.assertEqual(respuesta.data['archivado_por_nombre'], 'oficina')
        self.assertEqual(respuesta.data['planificacion']['estado'], 'ARCHIVADO')

        # Ficheros fuera...
        for ruta in (self.player, self.monitor, self.foto):
            self.assertFalse(ruta.exists(), ruta)
        self.assertFalse((Path(self.media) / 'imagenes' / str(self.project.pk)).exists())
        self.assertFalse(Imagen.objects.filter(modulo__proyecto=self.project).exists())
        self.assertFalse(FotoFabricacion.objects.filter(modulo__proyecto=self.project).exists())
        self.assertFalse(ProyectoCheckAdjunto.objects.filter(paso__proyecto=self.project).exists())
        self.project.refresh_from_db()
        self.assertFalse(self.project.plano_archivo)
        # ...y los datos de fabricacion dentro: las estadisticas no cambian.
        self.assertTrue(Modulo.objects.filter(pk=self.modulo.pk).exists())
        self.assertTrue(DetalleModuloFase.objects.filter(modulo=self.modulo).exists())
        self.assertTrue(MesaQueueItem.objects.filter(modulo=self.modulo, status='HECHO').exists())
        self.assertTrue(ProyectoCheck.objects.filter(proyecto=self.project).exists())
        self.assertEqual(self._modulos_completados(), 1)
        # Fuera del plan de fabricacion y no se puede volver a archivar ni descargar.
        self.assertNotIn(self.project.pk, plan_ferralla(self.factory.pk)['proyectos'])
        self.assertEqual(self.client.post(f'/api/proyectos/{self.project.pk}/archivar/', {'confirmado': True}, format='json').status_code, 400)
        self.assertEqual(self.client.post(f'/api/proyectos/{self.project.pk}/archivo/').status_code, 400)

    def test_no_se_archiva_lo_que_esta_en_produccion_ni_lo_pide_una_ferralla(self):
        linea = GrupoMesas.objects.create(nombre='Linea', usuario=self.factory)
        entrada = GrupoMesasProyecto.objects.create(grupo_mesas=linea, proyecto=self.project, orden=0)
        respuesta = self.client.post(f'/api/proyectos/{self.project.pk}/archivo/')
        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('cola', respuesta.data['detail'])
        entrada.delete()
        mesa = Mesa.objects.get(nombre='Mesa 1')
        MesaQueueItem.objects.create(mesa=mesa, modulo=self.modulo, fase='SUPERIOR', status='EN_COLA', position=1)
        respuesta = self.client.post(f'/api/proyectos/{self.project.pk}/archivo/')
        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('mesas', respuesta.data['detail'])

        self.client.force_authenticate(self.factory)
        self.assertEqual(self.client.post(f'/api/proyectos/{self.project.pk}/archivo/').status_code, 403)
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get('/api/proyectos/archivo-descarga/no-vale/').status_code, 403)
